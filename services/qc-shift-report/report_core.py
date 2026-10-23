"""Shift report calculations. No database/network access or business writes."""
from collections import defaultdict
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
import re
from zoneinfo import ZoneInfo

UTC = timezone.utc


def normalize_schedule(row):
    def clock(value):
        # MySQL TIME is returned as timedelta; older schemas store these as VARCHAR.
        if isinstance(value, timedelta):
            seconds = value.total_seconds()
            if not 0 <= seconds < 86400 or seconds != int(seconds):
                raise ValueError('Invalid configured shift time')
            return (datetime(2000, 1, 1) + value).time().isoformat()
        if not isinstance(value, str):
            raise ValueError('Invalid configured shift time')
        return datetime.strptime(value, '%H:%M:%S').time().isoformat()
    identity = str(row['id'])
    if not identity.isdigit() or int(identity) <= 0:
        raise ValueError('Invalid configured shift identity')
    next_day = row['nextDay']
    if type(next_day) not in (bool, int) or next_day not in (0, 1):
        raise ValueError('Invalid configured next-day flag')
    start, end = clock(row['start']), clock(row['end'])
    duration = (datetime.fromisoformat('2000-01-01T' + end) + timedelta(days=bool(next_day)) -
                datetime.fromisoformat('2000-01-01T' + start)).total_seconds()
    if not 0 < duration <= 86400:
        raise ValueError('Invalid configured shift duration')
    return dict(id=identity, start=start, end=end, nextDay=bool(next_day))


def shift_bounds(day, schedule, zone):
    parsed = date.fromisoformat(day)
    if parsed.isoformat() != day:
        raise ValueError('Invalid date')
    schedule = normalize_schedule(schedule)
    tz = ZoneInfo(zone)
    start = datetime.combine(parsed, time.fromisoformat(schedule['start']), tzinfo=tz)
    end = datetime.combine(parsed + timedelta(days=schedule['nextDay']),
                           time.fromisoformat(schedule['end']), tzinfo=tz)
    # Fail rather than silently move a configured boundary during a DST clock gap.
    for boundary in (start, end):
        if boundary.astimezone(UTC).astimezone(tz) != boundary:
            raise ValueError('Configured shift boundary does not exist on this date')
    return start.astimezone(UTC), end.astimezone(UTC)


def range_bounds(start_local, end_local, zone):
    """Interpret user-entered wall clocks in the warehouse zone, never the browser zone."""
    tz = ZoneInfo(zone)
    def parse(value):
        if not isinstance(value, str) or not re.fullmatch(r'\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(?::\d{2})?', value):
            raise ValueError('Use local dates and times without a timezone offset')
        local = datetime.fromisoformat(value)
        candidates = set()
        for fold in (0, 1):
            utc = local.replace(tzinfo=tz, fold=fold).astimezone(UTC)
            if utc.astimezone(tz).replace(tzinfo=None) == local:
                candidates.add(utc)
        if len(candidates) != 1:
            raise ValueError('The selected local time is missing or ambiguous because of daylight saving time')
        return local, candidates.pop()
    first, start = parse(start_local)
    last, end = parse(end_local)
    if not timedelta(0) < last - first <= timedelta(days=31) or end <= start:
        raise ValueError('End must follow start; select at most 31 days')
    return start, end


def report_period(report):
    zone = ZoneInfo(report['timezone'])
    start, end = [as_utc(report[key]).astimezone(zone) for key in ('startTime', 'endTime')]
    return f"{start:%Y-%m-%d %H:%M:%S %z} – {end:%Y-%m-%d %H:%M:%S %z} | {report['timezone']}"


def as_utc(value):
    if isinstance(value, str):
        value = datetime.fromisoformat(value.replace('Z', '+00:00'))
    if not isinstance(value, datetime):
        raise ValueError('Missing source timestamp')
    # MES auditing stores production, sample creation and modern QC timestamps in UTC.
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(UTC)


def inspector_username(value):
    if not isinstance(value, str) or not value.strip():
        return None
    value = value.strip()
    return None if value.casefold() in {'cwms-user', 'anonymous', 'anonymoususer'} else value


def attach_inspection_details(report, samples, inline, users):
    """Attach the exact counted records; registration is distinct from inspection."""
    groups = {(r['product'], r['machine']): r for r in report['rows']}
    start, end = as_utc(report['startTime']), as_utc(report['endTime'])
    for row in groups.values():
        row['firstSampleDetails'], row['inlineDetails'] = [], []
    for records, field in [(samples, 'firstSample'), (inline, 'inline')]:
        seen = {}
        for record in records:
            timestamp = as_utc(record['timestamp'])
            if record['warehouseId'] != report['warehouseId']:
                raise ValueError('Detail warehouse mismatch')
            if not start <= timestamp < end or (field == 'inline' and record.get('status') not in ('PASS', 'FAIL')):
                continue
            username = inspector_username(record.get('username'))
            person = users.get(username.casefold()) if username else None
            detail = dict(id=record['id'], number=record.get('number') or record['id'],
                          timestamp=timestamp.isoformat(), username=username,
                          displayName=person or username,
                          result='REGISTERED' if field == 'firstSample' else record['status'],
                          workOrderNumber=record.get('workOrderNumber') or '')
            if detail['id'] in seen:
                if seen[detail['id']] != detail:
                    raise ValueError('Conflicting inspection identity')
                continue
            seen[detail['id']] = detail
            groups[(record['product'], record['machine'])][field + 'Details'].append(detail)
        for row in groups.values():
            details = row[field + 'Details']
            details.sort(key=lambda r: (r['timestamp'], r['id']))
            if len(details) != row[field]:
                raise ValueError('Inspection details do not reconcile with count')
    return report


def aggregate(warehouse_id, day, schedule, zone, samples, inline, produced, finished, generated=None, *, bounds=None):
    if bounds is None:
        schedule = normalize_schedule(schedule)
        start, end = shift_bounds(day, schedule, zone)
    else:
        start, end = bounds
        schedule = None
    rows = defaultdict(lambda: dict(firstSample=0, inline=0, finishedGoods=0))
    warnings = []

    def eligible(record, completed=False):
        # Recheck scope/time even when the SQL adapter already applies them.
        if record.get('warehouseId') != warehouse_id:
            raise ValueError('Warehouse mismatch in report data')
        if completed and record.get('status') not in ('PASS', 'FAIL'):
            return False
        return start <= as_utc(record['timestamp']) < end

    def key(record):
        product, machine = record.get('product'), record.get('machine')
        if not isinstance(product, str) or not product.strip() or not isinstance(machine, str) or not machine.strip():
            raise ValueError('A source record cannot be mapped to a product and machine')
        return product, machine

    for records, field, completed in [(samples, 'firstSample', False), (inline, 'inline', True)]:
        seen = {}
        for record in records:
            if not eligible(record, completed):
                continue
            identity = record.get('id')
            if identity is None:
                raise ValueError('Missing source record identity')
            group = key(record)
            if identity in seen:
                if seen[identity] != group:
                    raise ValueError('Ambiguous inspection mapping')
                continue
            seen[identity] = group
            rows[group][field] += 1

    def pallets(records, completed=False):
        seen = {}
        for record in records:
            if not eligible(record, completed):
                continue
            lpn = record.get('lpn')
            if not isinstance(lpn, str) or not lpn.strip():
                raise ValueError('Missing production LPN')
            group = key(record)
            if lpn in seen and seen[lpn] != group:
                raise ValueError('An LPN maps to more than one product or machine')
            seen[lpn] = group
        return seen

    production = pallets(produced)
    inspections = pallets(finished, True)
    for group in production.values():
        rows[group]  # Include production with no inspections.
    for group in inspections.values():
        rows[group]['finishedGoods'] += 1
    for lpn in production.keys() & inspections.keys():
        if production[lpn] != inspections[lpn]:
            raise ValueError('Production and inspection LPN mappings disagree')
    covered = len(production.keys() & inspections.keys())
    previous = len(inspections.keys() - production.keys())
    if previous:
        warnings.append(f'{previous} inspected LPN(s) were not produced in this shift; excluded from coverage.')
    if end > datetime.now(UTC):
        warnings.append('This shift has not ended. Totals reflect records available at the generated time.')
    if bounds is not None:
        warnings = [message.replace('this shift', 'the selected period').replace('This shift', 'The selected period') for message in warnings]
    result_rows = [dict(product=p, machine=m, **counts) for (p, m), counts in sorted(rows.items())]
    return dict(warehouseId=warehouse_id, date=day, shift=schedule['id'] if schedule else 'custom', schedule=schedule, timezone=zone,
                startTime=start.isoformat(), endTime=end.isoformat(),
                generatedAt=(generated or datetime.now(UTC)).isoformat(), rows=result_rows,
                totals=dict(machines=len({r['machine'] for r in result_rows}),
                            **{field: sum(r[field] for r in result_rows) for field in ('firstSample', 'inline', 'finishedGoods')},
                            producedPallets=len(production), inspectedProducedPallets=covered,
                            inspectionRate=covered / len(production) if production else None), warnings=warnings)


def excel_bytes(report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'QC Shift Summary'
    def text(cell, value):
        cell.value = value
        cell.data_type = 's'  # User-controlled product/machine labels are never Excel formulas.
    text(ws['A1'], 'QC Shift Inspection Summary')
    ws.merge_cells('A1:E1')
    ws['A1'].font = Font(size=18, bold=True)
    text(ws['A2'], report_period(report))
    text(ws['A3'], 'Warehouse local time; end excluded')
    ws.merge_cells('A3:E3')
    headers = ['Product', 'Machine', '1st Sample', 'Inline', 'Finished Goods']
    for column, label in enumerate(headers, 1):
        cell = ws.cell(5, column, label)
        cell.fill = PatternFill('solid', fgColor='17324D')
        cell.font = Font(color='FFFFFF', bold=True)
    for index, record in enumerate(report['rows'], 6):
        text(ws.cell(index, 1), record['product'])
        text(ws.cell(index, 2), record['machine'])
        for column, field in enumerate(('firstSample', 'inline', 'finishedGoods'), 3):
            ws.cell(index, column, record[field]).number_format = '#,##0'
    last = 6 + len(report['rows'])
    ws.cell(last, 1, 'Total')
    ws.cell(last, 2, report['totals']['machines'])
    for column in range(3, 6):
        ws.cell(last, column, f'=SUM({get_column_letter(column)}6:{get_column_letter(column)}{last - 1})' if last > 6 else '=0')
    for cell in ws[last]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='EAF0F5')
    ws.cell(last + 2, 1, 'Pallets produced')
    ws.cell(last + 2, 5, report['totals']['producedPallets'])
    ws.cell(last + 3, 1, 'Produced pallets inspected')
    ws.cell(last + 3, 5, report['totals']['inspectedProducedPallets'])
    ws.cell(last + 4, 1, '% Inspected')
    ws.cell(last + 4, 5, f'=IF(E{last + 2}=0,"—",E{last + 3}/E{last + 2})').number_format = '0.00%'
    notes = ['1st Sample: sample registrations created in the shift. Editing an existing sample does not add another.',
             'Inline: completed PASS/FAIL work-order sampling inspections, including legacy result records.',
             'Finished Goods: distinct inspected production LPNs. Rechecks count once.',
             'Coverage: pallets produced AND inspected in this shift / pallets produced in this shift.',
             'Production counts distinct positive produced LPNs; fully reversed pallets are excluded.',
             f"Generated: {report['generatedAt']}"] + report['warnings']
    for index, note in enumerate(notes, last + 6):
        text(ws.cell(index, 1), note)
        ws.merge_cells(start_row=index, start_column=1, end_row=index, end_column=5)
        ws.cell(index, 1).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[index].height = 30
    for column, width in enumerate((38, 17, 18, 18, 22), 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = 'C6'
    ws.auto_filter.ref = f'A5:E{last - 1}' if last > 6 else 'A5:E5'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:5'
    ws.print_options.horizontalCentered = True
    ws.print_area = f'A1:E{last + 5 + len(notes)}'
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
