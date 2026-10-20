import io
import unittest
import base64,json,time
import threading
import urllib.request,urllib.error
from http.server import ThreadingHTTPServer
from datetime import datetime, timezone, timedelta
from unittest.mock import patch, Mock
from contextlib import contextmanager
from openpyxl import load_workbook
from report_core import aggregate, excel_bytes, shift_bounds
from report_pdf import pdf_bytes
from pypdf import PdfReader
from report_api import authorize, ReportError, positive, query_report
import report_api

UTC = timezone.utc
DAY = dict(id='2', start='06:00:00', end='18:00:00', nextDay=False)
NIGHT = dict(id='1', start='18:00:00', end='06:00:00', nextDay=True)


def record(identity='a', lpn='L1', timestamp='2026-10-18T12:00:00+00:00', **kwargs):
    return dict(id=identity, warehouseId=1, timestamp=timestamp, product='P1', machine='M1',
                lpn=lpn, status='PASS', **kwargs)


def report(samples=(), inline=(), produced=(), finished=()):
    return aggregate(1, '2026-10-18', DAY, 'America/Chicago', samples, inline, produced, finished)


class Calculations(unittest.TestCase):
    def test_shift_boundaries_and_dst(self):
        start, end = shift_bounds('2026-10-18', NIGHT, 'America/Chicago')
        self.assertEqual(start.isoformat(), '2026-10-18T23:00:00+00:00')
        self.assertEqual(end.isoformat(), '2026-10-19T11:00:00+00:00')
        start, end = shift_bounds('2026-10-31', NIGHT, 'America/Chicago')
        self.assertEqual((end-start).total_seconds(), 13*3600)
        start, end = shift_bounds('2026-03-07', NIGHT, 'America/Chicago')
        self.assertEqual((end-start).total_seconds(), 11*3600)
        with self.assertRaises(ValueError):
            shift_bounds('2026-02-30', DAY, 'UTC')
        with self.assertRaises(ValueError):
            shift_bounds('2026-10-18', dict(DAY,end='05:00:00'), 'UTC')

    def test_configured_times_arbitrary_shifts_and_invalid_configuration(self):
        early = dict(id='18', start='4:00:00', end='16:00:00', nextDay=0)
        start, end = shift_bounds('2026-10-18', early, 'America/Los_Angeles')
        self.assertEqual(start.isoformat(), '2026-10-18T11:00:00+00:00')
        self.assertEqual(end.isoformat(), '2026-10-18T23:00:00+00:00')
        third = dict(id='77', start='08:30:00', end='16:30:00', nextDay=False)
        self.assertEqual((shift_bounds('2026-10-18', third, 'UTC')[1] -
                          shift_bounds('2026-10-18', third, 'UTC')[0]).total_seconds(), 8*3600)
        for invalid in [dict(early, start=None), dict(early,nextDay=None), dict(early,end='25:00:00'),
                        dict(early,end='04:00:00'), dict(early,nextDay=True),dict(early,id='DAY')]:
            with self.subTest(invalid=invalid), self.assertRaises((ValueError,TypeError)):
                shift_bounds('2026-10-18',invalid,'UTC')

    def test_completed_records_and_half_open_interval(self):
        records = [record('start', timestamp='2026-10-18T11:00:00Z'),
                   record('end', timestamp='2026-10-18T23:00:00Z'),
                   record('before', timestamp='2026-10-18T10:59:59Z'), record('fail')]
        records[-1]['status'] = 'FAIL'
        for status in ['PENDING', 'SKIPPED', 'CANCELLED', 'NOT_APPLY']:
            row = record(status); row['status'] = status; records.append(row)
        value = report(inline=records)
        self.assertEqual(value['totals']['inline'], 2)

    def test_deduplication_prior_shift_and_same_product_multiple_machines(self):
        produced = [record(), record(), record('b', 'L2')]
        produced[-1]['machine'] = 'M2'
        checked = [record(), record('recheck'), record('old', 'L0')]
        value = report(samples=[record(), record()], inline=[record()], produced=produced, finished=checked)
        totals = value['totals']
        self.assertEqual(totals, dict(machines=2, firstSample=1, inline=1, finishedGoods=2,
                                     producedPallets=2, inspectedProducedPallets=1, inspectionRate=.5))
        self.assertEqual(len(value['rows']), 2)
        self.assertTrue(any('not produced' in warning for warning in value['warnings']))
        self.assertIsNone(report()['totals']['inspectionRate'])

    def test_missing_mapping_and_cross_warehouse_never_become_zero(self):
        for field, value in [('machine', None), ('product', ''), ('warehouseId', 2), ('timestamp', None), ('lpn', '')]:
            row = record(); row[field] = value
            with self.subTest(field=field), self.assertRaises(ValueError):
                report(produced=[row])
        duplicate = record(); duplicate['machine'] = 'M2'
        with self.assertRaises(ValueError):
            report(produced=[record(), duplicate])
        with self.assertRaises(ValueError):
            report(produced=[record()], finished=[duplicate])

    def test_attachment_control_totals(self):
        # Independent control totals from the supplied 35-machine attachment.
        samples = [dict(record(f's{i}'), machine=f'M{i % 35:02}') for i in range(36)]
        inline = [dict(record(f'q{i}'), machine=f'M{i % 35:02}') for i in range(179)]
        produced = [dict(record(f'p{i}', f'L{i}'), machine=f'M{i % 35:02}') for i in range(183)]
        finished = produced[:96]
        value = report(samples, inline, produced, finished)
        self.assertEqual(value['totals']['machines'], 35)
        self.assertEqual(value['totals']['firstSample'], 36)
        self.assertEqual(value['totals']['inline'], 179)
        self.assertEqual(value['totals']['finishedGoods'], 96)
        self.assertAlmostEqual(value['totals']['inspectionRate'], .5245901639344263)
        self.assertEqual(len(value['rows']), 35)

    def test_excel_uses_displayed_snapshot_and_safe_strings(self):
        row = record(); row['product'] = '=HYPERLINK("https://example.com")'
        value = report(samples=[row], produced=[row], finished=[row])
        wb = load_workbook(io.BytesIO(excel_bytes(value)), data_only=False)
        sheet = wb.active
        self.assertEqual(sheet['A6'].data_type, 's')
        self.assertEqual(sheet['A6'].value, row['product'])
        self.assertEqual(sheet['E6'].value, value['rows'][0]['finishedGoods'])
        self.assertEqual(sheet['E7'].value, '=SUM(E6:E6)')
        self.assertEqual(sheet['E11'].number_format, '0.00%')
        self.assertEqual(sheet.freeze_panes, 'C6')
        self.assertEqual(len(wb.sheetnames), 1)


class PdfReports(unittest.TestCase):
    def test_snapshot_text_empty_and_long_paginated_report(self):
        value = report()
        raw = pdf_bytes(value)
        self.assertTrue(raw.startswith(b'%PDF-'))
        reader = PdfReader(io.BytesIO(raw))
        self.assertEqual(len(reader.pages), 1)
        text = reader.pages[0].extract_text()
        self.assertIn('N/A', text)
        self.assertIn('No production or completed inspections', text)
        records = [dict(record(f's{i}'), product=f'PRODUCT-{i:03} ' + 'LONG PRODUCT LABEL ' * 5,
                        machine=f'M{i:03}') for i in range(95)]
        value = report(samples=records)
        value['rows'][0]['product'] = '中文产品 <b>literal & label</b>'
        raw = pdf_bytes(value)
        reader = PdfReader(io.BytesIO(raw))
        self.assertGreater(len(reader.pages), 2)
        text = '\n'.join(page.extract_text() for page in reader.pages)
        self.assertIn('中文产品', text)
        self.assertIn('<b>literal & label</b>', text)
        self.assertIn('PRODUCT-094', text)
        self.assertIn('Counting rules', text)
        for i, page in enumerate(reader.pages, 1):
            self.assertAlmostEqual(float(page.mediabox.width), 595.2756, places=2)
            self.assertIn(f'Page {i}', page.extract_text())
        self.assertEqual(value['totals']['firstSample'], 95)


class Access(unittest.TestCase):
    @patch.dict('os.environ', {'QC_SITE_TIMEZONES': '{"1":"America/Los_Angeles"}'})
    def test_empty_database_cursor_and_site_timezone_fallback(self):
        class Cursor:
            count=0
            def execute(self,*args): self.count+=1
            def fetchall(self):
                return ([{'company_id':1,'time_zone':None}] if self.count==1 else
                        [dict(DAY,start=timedelta(hours=4),end=timedelta(hours=16))] if self.count==2 else ())
        @contextmanager
        def empty_database(): yield Cursor()
        with patch('report_api.database',empty_database):
            value=query_report(1,1,'2026-10-18','2')
        self.assertEqual(value['rows'],[])
        self.assertIsNone(value['totals']['inspectionRate'])
        self.assertEqual(value['startTime'],'2026-10-18T11:00:00+00:00')

    def test_missing_config_and_removed_shift_do_not_fall_back(self):
        class Cursor:
            count=0
            def execute(self,*args): self.count+=1
            def fetchall(self):
                return ([dict(company_id=1,time_zone='UTC')] if self.count==1 else [])
        @contextmanager
        def db(): yield Cursor()
        with patch('report_api.database',db), self.assertRaises(ReportError) as error:
            query_report(1,1,'2026-10-18','2')
        self.assertEqual(error.exception.status,503)
        @contextmanager
        def populated(): yield Mock()
        with patch('report_api.database',populated), patch('report_api.load_shift_configuration',
                return_value=dict(timezone='UTC',shifts=[DAY])), self.assertRaises(ReportError) as error:
            query_report(1,1,'2026-10-18','999')
        self.assertEqual(error.exception.status,409)

    def test_invalid_scope(self):
        for value in ['0', '-1', '1 OR 1=1', None, True, 1.5]:
            with self.subTest(value=value), self.assertRaises(ReportError):
                positive(value)

    @patch.dict('os.environ', {'QC_GATEWAY_URL': 'http://gateway/api'})
    @patch('requests.get')
    def test_missing_token_company_and_qc_permissions(self, get):
        def token(company=1,username='QC',expiry=None,signature='mock-signature'):
            # Synthetic unit fixture only; the gateway is mocked in this test.
            payload=base64.urlsafe_b64encode(json.dumps(dict(companyId=company,sub=username,exp=expiry or time.time()+3600)).encode()).decode().rstrip('=')
            return 'Bearer e30.'+payload+'.'+signature
        response = Mock(status_code=200)
        get.return_value = response
        info = dict(user=dict(id=10, username='QC', companyId=1), menu=[dict(link='/qc/result')])
        response.json.return_value = dict(result=0, data=info)
        with self.assertRaises(ReportError):
            authorize('', 1, 1)
        header=token()
        owner = authorize(header, 1, 1)
        forwarded = get.call_args.kwargs['headers']
        self.assertEqual(forwarded, {'Authorization': header, 'username': 'QC', 'companyId': '1', 'warehouseId': '1'})
        self.assertNotEqual(owner, authorize(header, 1, 2))
        self.assertNotEqual(owner, authorize(token(signature='another-mock-signature'), 1, 1))
        with self.assertRaises(ReportError):
            authorize(header, 2, 1)
        with self.assertRaises(ReportError):
            authorize(token(username='OtherUser'),1,1)
        with self.assertRaises(ReportError):
            authorize(token(expiry=time.time()-1),1,1)
        with self.assertRaises(ReportError):
            authorize('Bearer not-a-token',1,1)
        info['menu'][0]['hide'] = True
        with self.assertRaises(ReportError):
            authorize(header, 1, 1)
        info['menu'] = []
        with self.assertRaises(ReportError):
            authorize(header, 1, 1)
        response.status_code = 401
        with self.assertRaises(ReportError):
            authorize(header, 1, 1)


class HttpContract(unittest.TestCase):
    def test_query_export_owner_expiry_and_untrusted_export_rows(self):
        value=report(samples=[record()],inline=[record()],produced=[record()],finished=[record()])
        def auth(header,*_):
            if not header: raise ReportError(401,'Sign in')
            return header
        server=ThreadingHTTPServer(('127.0.0.1',0),report_api.Handler)
        thread=threading.Thread(target=server.serve_forever,daemon=True)
        thread.start()
        base='http://127.0.0.1:'+str(server.server_port)+'/qc-shift-report/'
        def request(path,body=None,actor='Bearer fixture-a'):
            headers={'Content-Type':'application/json'}
            if actor: headers['Authorization']=actor
            return urllib.request.urlopen(urllib.request.Request(base+path,
                data=json.dumps(body).encode() if body is not None else None,headers=headers),timeout=3)
        path='report?warehouseId=1&companyId=1&date=2026-10-18&shift=2'
        try:
            with patch('report_api.authorize',side_effect=auth),patch('report_api.query_report',return_value=value):
                with self.assertRaises(urllib.error.HTTPError) as error:
                    request(path,actor=None)
                self.assertEqual(error.exception.code,401)
                with patch('report_api.query_shifts',return_value=dict(warehouseId=1,timezone='UTC',shifts=[DAY])) as shifts:
                    with request('shifts?warehouseId=1&companyId=1') as response:
                        self.assertEqual(json.load(response)['data']['shifts'][0]['id'],'2')
                    shifts.assert_called_once_with(1,1)
                with request(path) as response: saved=json.load(response)['data']
                body=dict(snapshotId=saved['snapshotId'],companyId=1,warehouseId=1)
                with request('export',body) as response:
                    self.assertEqual(response.headers['Cache-Control'],'no-store')
                    sheet=load_workbook(io.BytesIO(response.read())).active
                    self.assertEqual(sheet['D6'].value,saved['rows'][0]['inline'])
                with request('export-pdf',body) as response:
                    self.assertEqual(response.headers['Content-Type'], 'application/pdf')
                    self.assertIn('.pdf',response.headers['Content-Disposition'])
                    self.assertEqual(response.headers['Cache-Control'],'no-store')
                    text='\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(response.read())).pages)
                    self.assertIn('QC Shift Summary',text)
                    self.assertIn(saved['date'],text)
                    self.assertIn('P1',text)
                with self.assertRaises(urllib.error.HTTPError) as error:
                    request('export-pdf',body,actor='Bearer fixture-b')
                self.assertEqual(error.exception.code,409)
                with self.assertRaises(urllib.error.HTTPError) as error:
                    request('export-pdf',dict(body,rows=[{'firstSample':999}]))
                self.assertEqual(error.exception.code,400)
                with self.assertRaises(urllib.error.HTTPError) as error:
                    request('export',body,actor='Bearer fixture-b')
                self.assertEqual(error.exception.code,409)
                with self.assertRaises(urllib.error.HTTPError) as error:
                    request('export',dict(body,rows=[{'finishedGoods':999}]))
                self.assertEqual(error.exception.code,400)
                expiry,owner,data,size=report_api.CACHE[saved['snapshotId']]
                self.assertEqual(size,len(json.dumps(data).encode()))
                report_api.CACHE[saved['snapshotId']]=(0,owner,data,size)
                with self.assertRaises(urllib.error.HTTPError) as error:
                    request('export',body)
                self.assertEqual(error.exception.code,409)
                with self.assertRaises(urllib.error.HTTPError) as error:
                    request('export-pdf',body)
                self.assertEqual(error.exception.code,409)
        finally:
            server.shutdown();server.server_close();thread.join(timeout=2)
            report_api.CACHE.clear()


if __name__ == '__main__':
    unittest.main()
